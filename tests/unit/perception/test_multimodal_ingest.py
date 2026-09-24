"""InputRouter / Extractors / CertificateContentLoader 单元测试。"""

from __future__ import annotations

from pathlib import Path

import fitz
import pytest
from docx import Document

from label_agent.core.context import JobContext
from label_agent.core.exceptions import NonRetryableError
from label_agent.domain.models.content import OcrSegment
from label_agent.inference.providers.base import ModelCompletionResult
from label_agent.perception.certificate_content_loader import CertificateContentLoader
from label_agent.perception.extractors.base import AttachmentKind
from label_agent.perception.extractors.doc_extractor import DocExtractor
from label_agent.perception.extractors.docx_extractor import DocxExtractor
from label_agent.perception.extractors.pdf_extractor import PdfExtractor
from label_agent.perception.extractors.text_extractor import TextExtractor
from label_agent.perception.input_router import InputRouter
from label_agent.perception.text.normalizer import TextNormalizer


class _FakeGateway:
    """测试用 LLM 网关：图片请求返回固定描述文本。"""

    async def complete(self, request, model_alias=None):
        _ = model_alias
        return ModelCompletionResult(
            raw_text="视觉描述：证书含 UI 设计师字样",
            model_name="fake",
            prompt_tokens=1,
            completion_tokens=1,
            latency_ms=1,
        )


class _FakePromptManager:
    """测试用 PromptManager：返回固定 system/user。"""

    def render(self, scene: str, prompt_version: str, variables: dict) -> tuple[str, str]:
        _ = scene, prompt_version
        return "图片理解助手", f"hint={variables.get('context_hint')}"


def _ctx(payload: dict) -> JobContext:
    """构造最小 JobContext。"""
    return JobContext(job_id="j", task_id="t", scene="S1", biz_id="biz", payload=payload)


class _FakeCertificateOcr:
    version = "fake-certificate-ocr"

    async def recognize(self, image_uri: str, *, asset_id: str = "") -> list[OcrSegment]:
        _ = image_uri
        return [
            OcrSegment(
                asset_id=asset_id,
                text="软件工程师职业资格证书",
                bbox=[10, 10, 200, 10, 200, 40, 10, 40],
                confidence=0.95,
            )
        ]


def _loader() -> CertificateContentLoader:
    """装配带假 LLM 的 CertificateContentLoader。"""
    return CertificateContentLoader(
        text_normalizer=TextNormalizer(),
        media_guard=None,
        llm_gateway=_FakeGateway(),  # type: ignore[arg-type]
        prompt_manager=_FakePromptManager(),  # type: ignore[arg-type]
        vision_prompt_version="s1-vision-v1",
        scene="s1",
        ocr_engine=_FakeCertificateOcr(),  # type: ignore[arg-type]
    )


def test_input_router_legacy_and_attachments() -> None:
    """验证兼容字段与 attachments 协议均可路由，且按类型正确分类。"""
    router = InputRouter()
    refs = router.route(
        {
            "text": "补充说明",
            "attachments": [
                {"type": "pdf", "path": "/tmp/a.pdf"},
                {"type": "docx", "path": "/tmp/b.docx"},
                "/tmp/c.png",
            ],
        }
    )
    kinds = [r.kind for r in refs]
    assert AttachmentKind.TEXT in kinds
    assert AttachmentKind.PDF in kinds
    assert AttachmentKind.DOCX in kinds
    assert AttachmentKind.IMAGE in kinds
    assert router.version == "v2"


def test_input_router_input_path_suffix() -> None:
    """验证 input_path 按后缀推断 PDF。"""
    router = InputRouter()
    refs = router.route({"input_path": "D:/data/cert.pdf"})
    assert len(refs) == 1
    assert refs[0].kind == AttachmentKind.PDF


def test_text_extractor_inline() -> None:
    """验证内联文本抽取成功。"""
    from label_agent.perception.extractors.base import AttachmentRef

    result = TextExtractor().extract(
        AttachmentRef(kind=AttachmentKind.TEXT, ref="payload:text", inline_text="你好")
    )
    assert result.successful_modality == "text"
    assert result.text == "你好"


def test_docx_extractor(tmp_path: Path) -> None:
    """验证 DOCX 段落与表格可抽取。"""
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("职业：UI设计师")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "技能"
    table.rows[0].cells[1].text = "平面设计"
    doc.save(path)

    from label_agent.perception.extractors.base import AttachmentRef

    result = DocxExtractor().extract(
        AttachmentRef(kind=AttachmentKind.DOCX, ref=str(path), source_key="docx_path")
    )
    assert result.successful_modality == "docx"
    assert "UI设计师" in result.text
    assert "平面设计" in result.text


def test_doc_extractor_rejects_ole(tmp_path: Path) -> None:
    """验证原生 .doc（非 OOXML）返回失败模态并提示转码。"""
    path = tmp_path / "legacy.doc"
    # OLE 复合文档常见魔数
    path.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 32)

    from label_agent.perception.extractors.base import AttachmentRef

    result = DocExtractor().extract(
        AttachmentRef(kind=AttachmentKind.DOC, ref=str(path), source_key="input_path")
    )
    assert result.failed_modality == "doc"
    assert "docx" in (result.error_message or "").lower()


@pytest.mark.asyncio
async def test_pdf_text_extractor(tmp_path: Path) -> None:
    """验证数字 PDF 文本层抽取为 pdf_text。"""
    path = tmp_path / "cert.pdf"
    doc = fitz.open()
    page = doc.new_page()
    # 使用内置字体可稳定嵌入的 ASCII，避免 CJK 缺字导致文本层过短触发视觉降级
    page.insert_text(
        (72, 72),
        "Certificate holder is qualified as UI Designer professional career role",
    )
    doc.save(path)
    doc.close()

    from label_agent.perception.extractors.base import AttachmentRef

    result = await PdfExtractor().extract(
        AttachmentRef(kind=AttachmentKind.PDF, ref=str(path), source_key="pdf_path")
    )
    assert result.successful_modality == "pdf_text"
    assert "UI Designer" in result.text


@pytest.mark.asyncio
async def test_loader_scanned_pdf_uses_hybrid_ocr_and_vision(tmp_path: Path) -> None:
    """扫描 PDF 页面复用图片混合感知，并将临时图片 ID 映射为稳定页 ID。"""
    path = tmp_path / "scanned-certificate.pdf"
    doc = fitz.open()
    page = doc.new_page(width=320, height=200)
    page.draw_rect(fitz.Rect(20, 20, 300, 180), color=(0, 0, 0))
    doc.save(path)
    doc.close()

    bundle = await _loader().build(_ctx({"input_path": str(path)}))

    assert "软件工程师职业资格证书" in bundle.normalized_text
    assert "视觉描述" in bundle.normalized_text
    assert "pdf_hybrid" in bundle.successful_modalities
    assert "image_ocr" in bundle.successful_modalities
    assert "image_vision" in bundle.successful_modalities
    assert {segment.asset_id for segment in bundle.ocr_segments} == {"scanned-certificate-p1"}
    assert {block.asset_id for block in bundle.perception_text_blocks} == {"scanned-certificate-p1"}


@pytest.mark.asyncio
async def test_loader_text_only() -> None:
    """验证纯文本路径产出 ContentBundle。"""
    bundle = await _loader().build(_ctx({"text": "我是一名平面设计师"}))
    assert "平面设计师" in bundle.normalized_text
    assert "text" in bundle.successful_modalities
    assert bundle.perception_versions.get("input_router") == "v2"
    assert bundle.content_hash


@pytest.mark.asyncio
async def test_loader_image_uses_ocr_and_vision(tmp_path: Path) -> None:
    path = tmp_path / "certificate.png"
    pixmap = fitz.Pixmap(fitz.csRGB, fitz.IRect(0, 0, 320, 200), False)
    pixmap.clear_with(255)
    pixmap.save(str(path))

    bundle = await _loader().build(_ctx({"image_path": str(path)}))

    assert "软件工程师职业资格证书" in bundle.normalized_text
    assert "视觉描述" in bundle.normalized_text
    assert "image_ocr" in bundle.successful_modalities
    assert "image_vision" in bundle.successful_modalities
    assert len(bundle.ocr_segments) == 1
    assert {block.source_type for block in bundle.perception_text_blocks} == {
        "ocr",
        "vision",
    }


@pytest.mark.asyncio
async def test_loader_pdf(tmp_path: Path) -> None:
    """验证 CertificateContentLoader 可消费 PDF input_path。"""
    path = tmp_path / "s1.pdf"
    doc = fitz.open()
    page = doc.new_page()
    page.insert_text(
        (72, 100),
        "Certificate name: Packaging Designer Certification Material Document",
    )
    doc.save(path)
    doc.close()

    bundle = await _loader().build(_ctx({"input_path": str(path)}))
    assert "Packaging Designer" in bundle.normalized_text
    assert "pdf_text" in bundle.successful_modalities
    assert bundle.perception_versions.get("pdf_extractor") == "v1"


@pytest.mark.asyncio
async def test_loader_docx(tmp_path: Path) -> None:
    """验证 CertificateContentLoader 可消费 DOCX。"""
    path = tmp_path / "s1.docx"
    document = Document()
    document.add_paragraph("广告设计师相关材料")
    document.save(path)

    bundle = await _loader().build(_ctx({"docx_path": str(path)}))
    assert "广告设计师" in bundle.normalized_text
    assert "docx" in bundle.successful_modalities


@pytest.mark.asyncio
async def test_loader_empty_raises() -> None:
    """验证无附件时抛出不可重试错误。"""
    with pytest.raises(NonRetryableError):
        await _loader().build(_ctx({}))
